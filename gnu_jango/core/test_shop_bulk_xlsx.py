"""Excel (.xlsx) uploads and templates of the shop bulk tools (상품 일괄등록, 송장 일괄등록):
read with openpyxl in read-only mode, formulas never evaluated, numbers given back the way the CSV
path reads them, and damaged, oversized, zip-bomb and legacy .xls files refused."""

import datetime as dt
import io
import sys
import zipfile
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from gnu_jango.core.test_shop_bulk_tools import LABELS, SuperuserMixin, product_row
from gnu_jango.shop.models import ShopOrder, ShopOrderItem, ShopProduct
from gnu_jango.shop.product_import import FIELD_NAMES
from gnu_jango.shop.spreadsheet import SpreadsheetError, read_upload, xlsx_cell

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DELIVERY_HEADER = ["주문번호", "주문자명", "주문자전화1", "주문자전화2", "배송자명", "배송지전화1", "배송지전화2", "배송지주소", "배송회사", "운송장번호"]
ORDER_ID = "2026102100000000000001"


class Text(str):
    """A cell written as a string even when it starts with "=" (openpyxl would make a formula)."""


def xlsx_bytes(rows, *, cached=None):
    """An .xlsx made with openpyxl; ``cached`` maps a formula to the result Excel would have saved."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            if value in (None, ""):
                continue
            cell = sheet.cell(row=row_number, column=column, value=value)
            if isinstance(value, Text):
                cell.data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    for formula, result in (cached or {}).items():
        data = _rewrite_sheet(data, f"<f>{formula}</f><v />", f"<f>{formula}</f><v>{result}</v>")
    return data


def _rewrite_sheet(data, old, new):
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert old.encode() in content, old
                content = content.replace(old.encode(), new.encode())
            target.writestr(entry.filename, content)
    return output.getvalue()


def xlsx_upload(rows, *, name="upload.xlsx", cached=None):
    return SimpleUploadedFile(name, xlsx_bytes(rows, cached=cached), content_type=XLSX_TYPE)


def opened(response):
    return openpyxl.load_workbook(io.BytesIO(response.content)).worksheets[0]


class XlsxReadingTests(SimpleTestCase):
    def test_reads_the_first_sheet_as_text_like_the_csv_path(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["상품코드", "가격", "날짜"])
        workbook.active.append(["00123", 12000, dt.datetime(2026, 10, 22)])
        workbook.active.append([None, None, None])
        workbook.create_sheet("둘째").append(["읽지 않음"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        rows = read_upload(SimpleUploadedFile("a.xlsx", buffer.getvalue()), max_rows=10)
        self.assertEqual(rows, [["상품코드", "가격", "날짜"], ["00123", "12000", "2026-10-22"]])

    def test_cell_values_become_the_strings_the_importers_expect(self):
        self.assertEqual(xlsx_cell(12000.0), "12000")
        self.assertEqual(xlsx_cell(1234567890123.0), "1234567890123")
        self.assertEqual(xlsx_cell(2.0261021e19), "2.0261021E+19")  # digits already lost in Excel
        self.assertEqual(xlsx_cell(1.5), "1.5")
        self.assertEqual(xlsx_cell(True), "1")
        self.assertEqual(xlsx_cell(dt.datetime(2026, 1, 2, 3, 4, 5)), "2026-01-02 03:04:05")
        self.assertEqual(xlsx_cell(" '=SUM(A1) "), "=SUM(A1)")
        self.assertEqual(xlsx_cell("#N/A"), "")

    def test_formulas_are_not_evaluated(self):
        rows = read_upload(
            xlsx_upload([[Text("=1+1"), "=10000*2", "=2*3"]], cached={"10000*2": "777"}), max_rows=5
        )
        # a text cell stays text; a formula gives the value saved with it (777, not 20000), or nothing
        self.assertEqual(rows, [["=1+1", "777"]])

    def test_too_many_rows_and_far_away_cells_are_refused(self):
        with self.assertRaisesMessage(SpreadsheetError, "3행까지"):
            read_upload(xlsx_upload([["a"]] * 5), max_rows=3)
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "header"
        workbook.active["A200000"] = "stray"
        buffer = io.BytesIO()
        workbook.save(buffer)
        with self.assertRaisesMessage(SpreadsheetError, "빈 행 포함"):
            read_upload(SimpleUploadedFile("far.xlsx", buffer.getvalue()), max_rows=10)

    def test_damaged_and_fake_workbooks_are_refused(self):
        for name, data in (
            ("broken.xlsx", b"PK\x03\x04" + b"\x00" * 200),
            ("text.xlsx", b"just,a,csv\n"),
            ("half.xlsx", xlsx_bytes([["a"]])[:300]),
        ):
            with self.subTest(name=name), self.assertRaisesMessage(SpreadsheetError, "엑셀 파일을 읽을 수 없습니다"):
                read_upload(SimpleUploadedFile(name, data), max_rows=10)
        plain_zip = io.BytesIO()
        with zipfile.ZipFile(plain_zip, "w") as archive:
            archive.writestr("readme.txt", "not a workbook")
        with self.assertRaisesMessage(SpreadsheetError, "엑셀 파일을 읽을 수 없습니다"):
            read_upload(SimpleUploadedFile("zip.xlsx", plain_zip.getvalue()), max_rows=10)

    def test_a_zip_bomb_is_refused_before_it_is_unpacked(self):
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/worksheets/sheet1.xml", b"\x00" * (40 * 1024 * 1024))
        self.assertLess(len(bomb.getvalue()), 2 * 1024 * 1024)
        with self.assertRaisesMessage(SpreadsheetError, "압축을 푼 크기가 너무 큽니다"):
            read_upload(SimpleUploadedFile("bomb.xlsx", bomb.getvalue()), max_rows=10)

    def test_legacy_xls_and_other_workbooks_are_refused_with_a_way_out(self):
        ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
        with self.assertRaisesMessage(SpreadsheetError, "예전 엑셀 형식(.xls)"):
            read_upload(SimpleUploadedFile("old.xls", ole), max_rows=10)
        with self.assertRaisesMessage(SpreadsheetError, "예전 엑셀 형식(.xls)"):
            read_upload(SimpleUploadedFile("renamed.csv", ole), max_rows=10)
        with self.assertRaisesMessage(SpreadsheetError, "암호가 걸렸거나"):
            read_upload(SimpleUploadedFile("locked.xlsx", ole), max_rows=10)
        with self.assertRaisesMessage(SpreadsheetError, "이 형식은 읽을 수 없습니다"):
            read_upload(SimpleUploadedFile("macro.xlsm", xlsx_bytes([["a"]])), max_rows=10)
        with self.assertRaisesMessage(SpreadsheetError, "확장자를 .xlsx로"):
            read_upload(SimpleUploadedFile("workbook.csv", xlsx_bytes([["a"]])), max_rows=10)


class XlsxProductUploadTests(SuperuserMixin, TestCase):
    url = "admin:shop_shopproduct_bulk_upload"

    def post(self, upload, *, step="apply", mode="create"):
        return self.client.post(reverse(self.url), {"file": upload, "step": step, "mode": mode})

    def test_preview_saves_nothing_then_apply_creates_the_products(self):
        rows = [
            LABELS,
            list(FIELD_NAMES),
            product_row(it_id="X100", ca_id=10, it_name="셔츠", it_price=12000, it_stock_qty=5.0, it_type1=True),
            product_row(it_id=1700000002, ca_id="10", it_name="숫자 코드", it_price=500),
            product_row(it_id="X101", ca_id="99", it_name="없는 분류", it_price=1),
        ]
        response = self.post(xlsx_upload(rows), step="preview")
        self.assertContains(response, "미리보기")
        self.assertContains(response, "분류 &#x27;99&#x27;이(가) 없습니다.")
        self.assertFalse(ShopProduct.objects.exists())

        response = self.post(xlsx_upload(rows))
        self.assertContains(response, "반영")
        product = ShopProduct.objects.get(pk="X100")
        self.assertEqual((product.it_price, product.it_stock_qty, product.it_type1, product.category_id), (12000, 5, "1", "10"))
        self.assertTrue(ShopProduct.objects.filter(pk="1700000002", it_price=500).exists())
        self.assertFalse(ShopProduct.objects.filter(pk="X101").exists())

    def test_a_formula_looking_name_is_stored_as_text_and_a_formula_is_not_computed(self):
        rows = [
            LABELS,
            list(FIELD_NAMES),
            product_row(it_id="F100", ca_id="10", it_name=Text("=HYPERLINK(\"http://x\")"), it_price="=10000*2"),
        ]
        self.post(xlsx_upload(rows, cached={"10000*2": "777"}))
        product = ShopProduct.objects.get(pk="F100")
        self.assertEqual((product.it_name, product.it_price), ('=HYPERLINK("http://x")', 777))

    def test_legacy_xls_upload_shows_the_save_as_xlsx_message(self):
        response = self.post(SimpleUploadedFile("items.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64))
        self.assertContains(response, "Excel 통합 문서(.xlsx)")
        self.assertFalse(ShopProduct.objects.exists())

    def test_xlsx_template_opens_with_the_youngcart_columns_and_uploads_back(self):
        response = self.client.get(reverse(self.url) + "?download=template&format=xlsx")
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertIn("shop-product-upload-template.xlsx", response["Content-Disposition"])
        sheet = opened(response)
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header[:5], ["상품코드", "기본분류", "2차분류", "3차분류", "상품명"])
        self.assertEqual([cell.value for cell in sheet[2]], list(FIELD_NAMES))
        self.assertTrue(sheet["A1"].font.b)
        self.assertEqual((sheet["A3"].value, sheet["A3"].number_format), ("1700000001", "@"))
        response = self.post(SimpleUploadedFile("t.xlsx", response.content))
        self.assertTrue(ShopProduct.objects.filter(pk="1700000001", it_name="예시 상품").exists())

    def test_csv_template_is_still_the_default(self):
        response = self.client.get(reverse(self.url) + "?download=template")
        self.assertTrue(response["Content-Type"].startswith("text/csv"))

    def test_a_server_without_openpyxl_points_to_csv(self):
        upload = xlsx_upload([LABELS, list(FIELD_NAMES)])
        with patch.dict(sys.modules, {"openpyxl": None}):
            response = self.client.get(reverse(self.url) + "?download=template&format=xlsx", follow=True)
            self.assertContains(response, "openpyxl")
            response = self.post(upload)
        self.assertContains(response, "CSV 파일을 사용하세요")


class XlsxDeliveryUploadTests(SuperuserMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.url = reverse("admin:shop_shoporder_delivery_upload")
        self.ready = ShopOrder.objects.create(
            od_id=ORDER_ID, od_name="=HYPERLINK(\"http://evil.test\",\"x\")", od_email="buyer@example.test",
            od_b_name="김수령", od_status="준비", od_settle_case="무통장", od_cart_price=1000, od_receipt_price=1000,
        )
        ShopOrderItem.objects.create(order=self.ready, ct_id="L1", it_id="X", it_name="상품", ct_price=1000, ct_status="준비")

    def test_long_order_numbers_in_text_cells_match_and_numeric_ones_are_reported(self):
        blank = [""] * 7
        rows = [
            DELIVERY_HEADER,
            [ORDER_ID, *blank, "CJ대한통운", 123456789012],
            [2.0261021e19, *blank, "CJ대한통운", "1"],
            ["2026102100000000000009", *blank, "CJ대한통운", "#N/A"],
        ]
        response = self.client.post(self.url, {"file": xlsx_upload(rows), "step": "preview"})
        self.assertContains(response, "주문번호가 지수 형식")
        self.assertContains(response, "운송장번호가 없습니다")
        self.ready.refresh_from_db()
        self.assertEqual(self.ready.od_status, "준비")

        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(self.url, {"file": xlsx_upload(rows), "step": "apply"})
        self.ready.refresh_from_db()
        self.assertEqual((self.ready.od_status, self.ready.od_invoice), ("배송", "123456789012"))

    def test_xlsx_template_keeps_order_numbers_as_text_and_names_out_of_formulas(self):
        response = self.client.get(self.url + "?download=template&format=xlsx")
        self.assertIn("shop-delivery-upload.xlsx", response["Content-Disposition"])
        sheet = opened(response)
        self.assertEqual([cell.value for cell in sheet[1]], DELIVERY_HEADER)
        order_cell = sheet["A2"]
        self.assertEqual((order_cell.value, order_cell.data_type, order_cell.number_format), (ORDER_ID, "s", "@"))
        self.assertEqual(sheet["J2"].number_format, "@")
        self.assertEqual((sheet["B2"].value, sheet["B2"].data_type), (self.ready.od_name, "s"))
        self.assertEqual(sheet.freeze_panes, "A2")
