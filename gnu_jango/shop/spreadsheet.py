"""Reading and writing the spreadsheet files the bulk shop tools exchange with spreadsheet programs.

youngcart reads .xls uploads with PHPExcel (adm/shop_admin/itemexcelupdate.php,
orderdeliveryupdate.php). Here the tools take an Excel workbook (.xlsx, read with openpyxl) or CSV
("CSV UTF-8" in Excel); the legacy binary .xls is refused. Nothing in an upload is evaluated: an
.xlsx formula cell gives the value Excel saved with it and is never calculated here, and a text
cell is only text. A cell we write that starts like a formula is kept from becoming one (``'`` in
CSV, a string cell in .xlsx); the CSV prefix is undone on reading, so a downloaded template
uploads back unchanged.
"""

import csv
import datetime as dt
import io
import logging
import os
import re
import warnings
import zipfile
from decimal import Decimal

from django.http import HttpResponse

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
# Excel saves "CSV UTF-8" with a BOM; an older "CSV" save of a Korean sheet is CP949
ENCODINGS = ("utf-8-sig", "cp949")

XLSX_EXTENSIONS = (".xlsx",)
LEGACY_EXTENSIONS = (".xls",)
# workbooks openpyxl cannot read, or that we do not take (.xlsm carries macros)
UNSUPPORTED_EXTENSIONS = (".xlsm", ".xlsb", ".ods", ".numbers")
ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .xls, and a password-protected .xlsx
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# A 2MB .xlsx unpacks to well under this; a zip bomb declares far more. zipfile never inflates an
# entry past its declared size, so the declared total bounds what openpyxl can unpack.
MAX_XLSX_UNPACKED_BYTES = 30 * 1024 * 1024
MAX_XLSX_ENTRIES = 300
# columns read from each row (the product sheet has 40); cells further right are not read
MAX_XLSX_COLUMNS = 100
# blank rows tolerated on top of the row limit before the sheet is refused: openpyxl yields every
# row up to the last one used, so one stray cell in row 1,048,576 must not mean a million rows
BLANK_ROW_ALLOWANCE = 5000
# Excel keeps 15 significant digits: an integer below this read from a number cell is exact
EXCEL_EXACT_INTEGER_LIMIT = 10**15
XLSX_ERROR_VALUES = frozenset({"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"})
XLSX_HEADER_FILL = "FFE5E7EB"
XLSX_COLUMN_WIDTH = 16
XLSX_TEXT_COLUMN_WIDTH = 24

TEXT_CODE_RE = re.compile(r"[0-9A-Za-z_-]{1,64}")
EXCEL_TEXT_RE = re.compile(r'="([0-9A-Za-z_-]{1,64})"')

LEGACY_XLS_MESSAGE = (
    "예전 엑셀 형식(.xls)은 읽을 수 없습니다. 엑셀에서 '다른 이름으로 저장 → Excel 통합 문서(.xlsx)' "
    "또는 'CSV UTF-8(쉼표로 분리)'로 저장해 올려 주세요."
)
PROTECTED_XLSX_MESSAGE = (
    "암호가 걸렸거나 예전 형식(.xls)인 파일입니다. 암호를 풀고 'Excel 통합 문서(.xlsx)'로 저장해 올려 주세요."
)
UNSUPPORTED_MESSAGE = "이 형식은 읽을 수 없습니다. 'Excel 통합 문서(.xlsx)' 또는 'CSV UTF-8(쉼표로 분리)'로 저장해 올려 주세요."
BROKEN_XLSX_MESSAGE = "엑셀 파일을 읽을 수 없습니다. 파일이 손상되었을 수 있습니다. 엑셀에서 다시 저장해 올려 주세요."
NOT_CSV_MESSAGE = (
    "CSV 파일이 아닙니다. 엑셀 파일이면 확장자를 .xlsx로 두고, CSV라면 'CSV UTF-8(쉼표로 분리)' 형식으로 저장해 올려 주세요."
)
NO_OPENPYXL_MESSAGE = "이 서버에는 엑셀(.xlsx) 처리 모듈(openpyxl)이 설치되어 있지 않습니다. CSV 파일을 사용하세요."


class SpreadsheetError(ValueError):
    """The upload cannot be read at all (the message is shown to the admin)."""


class ExcelText(str):
    """A code Excel must keep as text: a 20-digit order number opened from a CSV would become
    2.03E+19 and lose its digits. Written as ="..." in CSV (Excel shows, and saves back, the text)
    and as a text-formatted string cell in .xlsx."""


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (int, Decimal)) and not isinstance(value, bool):
        return str(value)  # a number (a negative refund total) is not a formula
    text = str(value)
    if isinstance(value, ExcelText) and TEXT_CODE_RE.fullmatch(text):
        return f'="{text}"'
    return "'" + text if text[:1] in FORMULA_PREFIXES else text


def read_cell(value):
    text = str(value or "").strip()
    wrapped = EXCEL_TEXT_RE.fullmatch(text)
    if wrapped:
        return wrapped.group(1)
    # a cell this module wrote as '=..., or one typed that way to keep Excel from calculating
    if len(text) > 1 and text[0] == "'" and text[1] in FORMULA_PREFIXES:
        return text[1:]
    return text


def _too_many_rows(max_rows):
    return SpreadsheetError(f"한 번에 {max_rows:,}행까지 올릴 수 있습니다. 파일을 나눠 올려 주세요.")


def _too_large(max_bytes):
    return SpreadsheetError(f"파일이 너무 큽니다. {max_bytes // (1024 * 1024)}MB 이하로 나눠 올려 주세요.")


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # requirements.txt installs it; an older deployment may lack it
        raise SpreadsheetError(NO_OPENPYXL_MESSAGE) from exc
    return openpyxl


def read_upload(uploaded, *, max_rows, max_bytes=MAX_UPLOAD_BYTES):
    """The rows of an uploaded .xlsx (first sheet) or CSV file as lists of stripped strings, blank
    rows dropped. The format follows the extension, checked against the file's first bytes.

    Raises SpreadsheetError for a missing, oversized, legacy .xls, damaged or unreadable file, or
    one with more than ``max_rows`` non-blank rows (header rows included in the count)."""
    if not uploaded:
        raise SpreadsheetError("업로드할 파일(.xlsx 또는 .csv)을 선택하세요.")
    extension = os.path.splitext(str(getattr(uploaded, "name", "") or "").lower())[1]
    size = getattr(uploaded, "size", None)
    if size is not None and size > max_bytes:
        raise _too_large(max_bytes)
    raw = uploaded.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise _too_large(max_bytes)
    if raw[:8] == OLE_MAGIC:
        raise SpreadsheetError(PROTECTED_XLSX_MESSAGE if extension in XLSX_EXTENSIONS else LEGACY_XLS_MESSAGE)
    if extension in LEGACY_EXTENSIONS:
        raise SpreadsheetError(LEGACY_XLS_MESSAGE)
    if extension in UNSUPPORTED_EXTENSIONS:
        raise SpreadsheetError(UNSUPPORTED_MESSAGE)
    if extension in XLSX_EXTENSIONS:
        if raw[:4] != ZIP_MAGIC:
            raise SpreadsheetError(BROKEN_XLSX_MESSAGE)
        return _read_xlsx(raw, max_rows=max_rows)
    return _read_csv(raw, max_rows=max_rows)


def _read_csv(raw, *, max_rows):
    if raw[:2] == b"PK" or b"\x00" in raw[:4096]:
        raise SpreadsheetError(NOT_CSV_MESSAGE)
    text = None
    for encoding in ENCODINGS:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SpreadsheetError("파일의 문자 인코딩을 읽을 수 없습니다. UTF-8로 저장해 올려 주세요.")
    rows = []
    try:
        for row in csv.reader(io.StringIO(text, newline="")):
            cells = [read_cell(cell) for cell in row]
            if not any(cells):
                continue
            rows.append(cells)
            if len(rows) > max_rows:
                raise _too_many_rows(max_rows)
    except csv.Error as exc:
        raise SpreadsheetError(f"CSV 형식을 읽을 수 없습니다: {exc}") from exc
    if not rows:
        raise SpreadsheetError("파일에 내용이 없습니다.")
    return rows


def _check_xlsx_archive(raw):
    """Refuse a damaged zip, or one that would unpack too large, before openpyxl opens it."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, OSError) as exc:
        raise SpreadsheetError(BROKEN_XLSX_MESSAGE) from exc
    if len(entries) > MAX_XLSX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_XLSX_UNPACKED_BYTES:
        raise SpreadsheetError(
            "엑셀 파일의 압축을 푼 크기가 너무 큽니다. 내용만 새 통합 문서에 붙여 넣어 저장하거나 파일을 나눠 올려 주세요."
        )
    names = {entry.filename for entry in entries}
    if any(entry.flag_bits & 0x1 for entry in entries) or "[Content_Types].xml" not in names:
        raise SpreadsheetError(BROKEN_XLSX_MESSAGE)


def _read_xlsx(raw, *, max_rows):
    openpyxl = _openpyxl()
    _check_xlsx_archive(raw)
    try:
        with warnings.catch_warnings():
            # openpyxl warns about parts it skips (data validation extensions, default styles)
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure of an uploaded file is "unreadable"
        logger.info("Unreadable .xlsx upload: %r", exc)
        raise SpreadsheetError(BROKEN_XLSX_MESSAGE) from exc
    try:
        return _xlsx_rows(workbook, max_rows=max_rows)
    except SpreadsheetError:
        raise
    except Exception as exc:  # noqa: BLE001 - the sheet XML is parsed lazily, row by row
        logger.info("Unreadable .xlsx sheet: %r", exc)
        raise SpreadsheetError(BROKEN_XLSX_MESSAGE) from exc
    finally:
        workbook.close()


def _xlsx_rows(workbook, *, max_rows):
    if not workbook.worksheets:
        raise SpreadsheetError("엑셀 파일에 시트가 없습니다.")
    sheet = workbook.worksheets[0]
    scan_limit = max_rows + BLANK_ROW_ALLOWANCE
    rows = []
    for scanned, values in enumerate(sheet.iter_rows(max_col=MAX_XLSX_COLUMNS, values_only=True), start=1):
        if scanned > scan_limit:
            raise SpreadsheetError(
                f"첫 번째 시트의 행이 너무 많습니다 (빈 행 포함 {scan_limit:,}행 초과). "
                "내용 아래의 빈 행을 삭제하거나 파일을 나눠 올려 주세요."
            )
        cells = [xlsx_cell(value) for value in values]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        rows.append(cells)
        if len(rows) > max_rows:
            raise _too_many_rows(max_rows)
    if not rows:
        raise SpreadsheetError("첫 번째 시트에 내용이 없습니다.")
    return rows


def xlsx_cell(value):
    """An .xlsx cell value as the text the CSV path would give for the same sheet."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < EXCEL_EXACT_INTEGER_LIMIT:
            return str(int(value))  # 12000.0 -> "12000", a 13-digit code stays whole
        # a longer number already lost digits in Excel: give the 2.03E+19 form Excel shows, which
        # the importers report (an order number) or expand like youngcart (a product code)
        return format(value, ".15g").upper()
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d") if value.time() == dt.time() else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    text = read_cell(value)
    # an error a formula left (=VLOOKUP(...) -> #N/A) is no value: it must not become an invoice number
    return "" if text in XLSX_ERROR_VALUES else text


def csv_download(filename, header_rows, rows):
    """A CSV attachment Excel opens as UTF-8 (BOM), with formula-like cells neutralised."""
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("﻿")
    writer = csv.writer(response)
    for row in (*header_rows, *rows):
        writer.writerow([safe_cell(value) for value in row])
    return response


def xlsx_download(filename, header_rows, rows, *, text_columns=()):
    """An .xlsx attachment: header rows shaded and frozen, every text a string cell (never a
    formula), and the ``text_columns`` (0-based) formatted as text so a long order number or
    invoice typed into them is not turned into 2.03E+19. Raises SpreadsheetError without openpyxl."""
    openpyxl = _openpyxl()
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "양식"
    all_rows = (*header_rows, *rows)
    for index in range(max((len(row) for row in all_rows), default=0)):
        dimension = sheet.column_dimensions[get_column_letter(index + 1)]
        dimension.width = XLSX_TEXT_COLUMN_WIDTH if index in text_columns else XLSX_COLUMN_WIDTH
        if index in text_columns:
            dimension.number_format = "@"
    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", start_color=XLSX_HEADER_FILL, end_color=XLSX_HEADER_FILL)
    for row_number, row in enumerate(all_rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=column)
            if row_number <= len(header_rows):
                cell.font = header_font
                cell.fill = header_fill
            if column - 1 in text_columns or isinstance(value, ExcelText):
                cell.number_format = "@"
            _write_xlsx_value(cell, value, ILLEGAL_CHARACTERS_RE)
    if header_rows:
        sheet.freeze_panes = sheet.cell(row=len(header_rows) + 1, column=1)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _write_xlsx_value(cell, value, illegal_characters):
    if value is None or value == "":
        return
    if isinstance(value, (int, Decimal)) and not isinstance(value, bool):
        cell.value = value
        return
    # control characters openpyxl refuses to write (an address pasted from another program)
    cell.value = illegal_characters.sub("", str(value))
    cell.data_type = "s"  # a value starting with "=" stays text, not a formula
